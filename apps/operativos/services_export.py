import csv
import io
from datetime import datetime

from django.http import HttpResponse

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.units import mm
from reportlab.lib.colors import HexColor
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER

from .models import Operativo

COLOR_PRIMARIO = HexColor('#6C5CE7')
COLOR_HEADER = HexColor('#6C5CE7')
COLOR_HEADER_TEXT = HexColor('#FFFFFF')


def _operativo_alumnos_qs(operativo: Operativo):
    return operativo.alumnos.select_related('curso', 'paciente').order_by('apellido', 'nombre')


def _alumno_row_dict(alumno):
    try:
        med = alumno.evaluacion_medica
        med_ok = bool(med.completada)
        med_peso = med.peso
        med_talla = med.talla
        med_imc = med.imc
        med_pas = med.pas
        med_pad = med.pad
        med_trajo = med.trajo_carnet
        med_completo = med.carnet_completo
    except Exception:
        med_ok = False
        med_peso = med_talla = med_imc = med_pas = med_pad = None
        med_trajo = med_completo = False

    try:
        odonto = alumno.evaluacion_odontologica
        odonto_ok = bool(odonto.completada)
        odonto_salud = odonto.salud_bucal
        odonto_cpo = f"{odonto.cpo_c or 0}/{odonto.cpo_p or 0}/{odonto.cpo_o or 0}"
        odonto_ceo = f"{odonto.ceo_c or 0}/{odonto.ceo_e or 0}/{odonto.ceo_o or 0}"
    except Exception:
        odonto_ok = False
        odonto_salud = ''
        odonto_cpo = odonto_ceo = ''

    curso_label = ''
    try:
        if alumno.curso:
            curso_label = f"{alumno.curso.sala_grado_anio or ''} {alumno.curso.division or ''}".strip()
    except Exception:
        pass

    return {
        'dni': alumno.dni,
        'apellido': alumno.apellido,
        'nombre': alumno.nombre,
        'sexo': alumno.sexo or '',
        'fecha_nacimiento': alumno.fecha_nacimiento.isoformat() if alumno.fecha_nacimiento else '',
        'curso': curso_label or 'Sin curso',
        'estado': alumno.estado,
        'completo': bool(alumno.completo),
        'escuela_completado': bool(alumno.escuela_completado),
        'medica_ok': med_ok,
        'med_peso': str(med_peso) if med_peso is not None else '',
        'med_talla': str(med_talla) if med_talla is not None else '',
        'med_imc': str(med_imc) if med_imc is not None else '',
        'med_pas_pad': f"{med_pas or ''}/{med_pad or ''}".strip('/') if (med_pas or med_pad) else '',
        'med_carnet': f"{'Sí' if med_trajo else 'No'}/{'Sí' if med_completo else 'No'}",
        'odonto_ok': odonto_ok,
        'odonto_salud': odonto_salud or '',
        'odonto_cpo': odonto_cpo,
        'odonto_ceo': odonto_ceo,
    }


def generar_export_csv(operativo: Operativo) -> bytes:
    alumnos = _operativo_alumnos_qs(operativo)
    output = io.StringIO()
    writer = csv.writer(output, delimiter=',', quotechar='"', quoting=csv.QUOTE_MINIMAL)
    headers = ['dni','apellido','nombre','sexo','fecha_nacimiento','curso','estado','completo','escuela_completado','medica_ok','peso','talla','imc','pas_pad','carnet','odonto_ok','odonto_salud','cpo','ceo']
    writer.writerow(headers)
    for a in alumnos:
        d = _alumno_row_dict(a)
        writer.writerow([d['dni'], d['apellido'], d['nombre'], d['sexo'], d['fecha_nacimiento'], d['curso'], d['estado'], d['completo'], d['escuela_completado'], d['medica_ok'], d['med_peso'], d['med_talla'], d['med_imc'], d['med_pas_pad'], d['med_carnet'], d['odonto_ok'], d['odonto_salud'], d['odonto_cpo'], d['odonto_ceo']])
    return output.getvalue().encode('utf-8-sig')


def generar_export_excel(operativo: Operativo) -> bytes:
    alumnos = list(_operativo_alumnos_qs(operativo))
    wb = Workbook()
    ws = wb.active
    ws.title = 'Alumnos'
    headers = ['DNI','Apellido','Nombre','Sexo','Fecha nac.','Curso','Estado','Completo','Escuela','Médica','Peso','Talla','IMC','PAS/PAD','Carnet','Odonto','Salud bucal','CPO','ceo']
    ws.append(headers)
    # Style header
    header_fill = PatternFill(start_color='6C5CE7', end_color='6C5CE7', fill_type='solid')
    header_font = Font(color='FFFFFF', bold=True, size=9)
    thin = Side(style='thin', color='DFE6E9')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = border
    # Rows
    for a in alumnos:
        d = _alumno_row_dict(a)
        ws.append([d['dni'], d['apellido'], d['nombre'], d['sexo'], d['fecha_nacimiento'], d['curso'], d['estado'], 'Sí' if d['completo'] else 'No', 'Sí' if d['escuela_completado'] else 'No', 'Sí' if d['medica_ok'] else 'No', d['med_peso'], d['med_talla'], d['med_imc'], d['med_pas_pad'], d['med_carnet'], 'Sí' if d['odonto_ok'] else 'No', d['odonto_salud'], d['odonto_cpo'], d['odonto_ceo']])
    # Borders for data
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=len(headers)):
        for cell in row:
            cell.border = border
            cell.font = Font(size=9)
            cell.alignment = Alignment(vertical='center')
    # Column widths
    widths = [14,16,16,8,13,14,12,10,10,10,8,8,8,10,12,10,14,10,10]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
    ws.freeze_panes = 'A2'
    # Info sheet (operativo)
    ws2 = wb.create_sheet('Operativo')
    ws2.append(['Campo','Valor'])
    for c in ws2[1]:
        c.fill = header_fill
        c.font = header_font
    ws2.append(['Operativo', operativo.nombre or ''])
    ws2.append(['Escuela', getattr(operativo.escuela, 'nombre', '')])
    ws2.append(['CUE', getattr(operativo.escuela, 'cue', '') or ''])
    ws2.append(['Fecha', str(operativo.fecha)])
    ws2.append(['Lugar', operativo.get_lugar_realizacion_display()])
    ws2.append(['Estado', operativo.get_estado_display()])
    ws2.append(['Total alumnos', operativo.alumnos.count()])
    for row in ws2.iter_rows(min_row=2, max_row=ws2.max_row, max_col=2):
        for c in row:
            c.border = border
            c.font = Font(size=9)
    ws2.column_dimensions['A'].width = 18
    ws2.column_dimensions['B'].width = 40

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def generar_export_pdf(operativo: Operativo) -> bytes:
    alumnos = list(_operativo_alumnos_qs(operativo))
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        leftMargin=10 * mm,
        rightMargin=10 * mm,
        topMargin=12 * mm,
        bottomMargin=10 * mm,
        title=f'PROSANE Export - {operativo.id}',
    )
    styles = getSampleStyleSheet()
    s_title = ParagraphStyle('TitleCustom', parent=styles['Title'], textColor=COLOR_PRIMARIO, fontSize=16, alignment=TA_CENTER, spaceAfter=2)
    s_sub = ParagraphStyle('Sub', parent=styles['Normal'], fontSize=8, textColor=HexColor('#636E72'), alignment=TA_CENTER, spaceAfter=6)
    s_cell = ParagraphStyle('Cell', parent=styles['Normal'], fontSize=6, leading=7)
    s_cell_bold = ParagraphStyle('CellBold', parent=s_cell, fontName='Helvetica-Bold', textColor=COLOR_HEADER_TEXT)
    story = []
    story.append(Paragraph('PROSANE — Resumen de operativo', s_title))
    escuela_nombre = getattr(operativo.escuela, 'nombre', '')
    story.append(Paragraph(f"{_safe(escuela_nombre)} &nbsp;|&nbsp; {operativo.fecha} &nbsp;|&nbsp; {operativo.get_lugar_realizacion_display()} &nbsp;|&nbsp; {operativo.get_estado_display()} &nbsp;|&nbsp; {len(alumnos)} alumnos", s_sub))
    story.append(HRFlowable(width='100%', thickness=1, color=COLOR_PRIMARIO, spaceAfter=6))

    headers = ['DNI','Apellido','Nombre','Curso','Estado','Comp.','Méd.','Od.','Esc.']
    # Build table data
    data = [[Paragraph(f"<b>{h}</b>", s_cell_bold) for h in headers]]
    for a in alumnos:
        d = _alumno_row_dict(a)
        row = [
            Paragraph(d['dni'], s_cell),
            Paragraph(d['apellido'], s_cell),
            Paragraph(d['nombre'], s_cell),
            Paragraph(d['curso'], s_cell),
            Paragraph(d['estado'], s_cell),
            Paragraph('Sí' if d['completo'] else 'No', s_cell),
            Paragraph('Sí' if d['medica_ok'] else 'No', s_cell),
            Paragraph('Sí' if d['odonto_ok'] else 'No', s_cell),
            Paragraph('Sí' if d['escuela_completado'] else 'No', s_cell),
        ]
        data.append(row)

    # colWidths landscape: total ~277mm usable (297 -20)
    col_widths = [26*mm, 28*mm, 28*mm, 24*mm, 22*mm, 16*mm, 16*mm, 16*mm, 16*mm]
    t = Table(data, colWidths=col_widths, repeatRows=1)
    t.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), COLOR_HEADER),
        ('TEXTCOLOR', (0, 0), (-1, 0), COLOR_HEADER_TEXT),
        ('GRID', (0, 0), (-1, -1), 0.4, HexColor('#DFE6E9')),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('LEFTPADDING', (0, 0), (-1, -1), 3),
        ('RIGHTPADDING', (0, 0), (-1, -1), 3),
        ('TOPPADDING', (0, 0), (-1, -1), 3),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [HexColor('#FFFFFF'), HexColor('#F8F9FA')]),
    ]))
    story.append(t)
    story.append(Spacer(1, 8))
    from django.utils import timezone
    ahora = timezone.now().strftime('%d/%m/%Y %H:%M')
    s_small = ParagraphStyle('Small', parent=styles['Normal'], fontSize=6, textColor=HexColor('#636E72'), alignment=TA_CENTER)
    story.append(Paragraph(f"Generado el {ahora} — Operativo {operativo.id} — PROSANE Salta", s_small))
    doc.build(story)
    pdf = buffer.getvalue()
    buffer.close()
    return pdf


def _safe(v, default='—'):
    if v is None or v == '':
        return default
    return str(v)
